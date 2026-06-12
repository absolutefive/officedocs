"""Deck 모델 + PPTX 템플릿 → 최종 PPTX 빌더.

모든 내용은 자리표시자/텍스트 상자/표/그림 등 네이티브 PowerPoint 개체로
생성되므로, 결과물을 PowerPoint에서 직접 수정할 수 있다.

두 가지 템플릿 방식을 지원한다.

- **레이아웃 방식**: 슬라이드 레이아웃의 자리표시자에 내용을 채운다.
  (표준 Office 계열 템플릿)
- **프로토타입 방식**: 디자인이 예시 슬라이드의 텍스트 상자에 들어있는
  기업 템플릿을 위해, 예시 슬라이드를 복제한 뒤 마커 텍스트를 치환한다.
  (사이드카 ``*.layouts.json`` 으로 정의 — layouts.py 참고)
"""

from __future__ import annotations

import copy
from pathlib import Path
from typing import Dict, List, Optional

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from officedocs.layouts import LayoutResolver, PrototypeSpec
from officedocs.model import (
    Bar,
    Block,
    Card,
    CardGrid,
    CodeBlock,
    Deck,
    Image,
    Paragraph,
    Run,
    Slide,
    Table,
)

_TITLE_TYPES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
_BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
_CODE_FONT = "Consolas"
_TEXTBOX_FONT_SIZE = Pt(14)
_INDENT_PER_LEVEL = Emu(228600)  # 0.25in

# 카드/강조 바 디자인 토큰 (gen_deck.js에서 이식). 포인트 컬러는 템플릿
# 테마의 강조1을 따라가므로 어떤 템플릿에서든 어울리게 렌더링된다.
_CARD_GRAY = RGBColor.from_string("F4F4F5")
_CARD_LINE = RGBColor.from_string("E9E9EC")
_CARD_SUB = RGBColor.from_string("6B7077")
_CARD_MUTED = RGBColor.from_string("8A8F98")
_ARROW_GRAY = RGBColor.from_string("9AA0A6")
_DISP_FONT = "Pretendard ExtraBold"

# 다크 코드 윈도우 디자인 토큰 (AgentOS 스타일, gen_deck.js에서 이식)
_CODE_BG = RGBColor.from_string("1F2335")
_CODE_HEAD = RGBColor.from_string("171A2B")
_CODE_TXT = RGBColor.from_string("E8EAF2")
_CODE_DIM = RGBColor.from_string("8B90A5")
_CODE_DOTS = ("ED6A5E", "F5BE4F", "61C554")  # 신호등 점: 빨강/노랑/초록
_CODE_MONO = "D2Coding"

# 프로토타입 슬라이드에서 복제 대상이 되는 도형 요소
_CLONABLE_TAGS = tuple(
    qn(t) for t in ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp")
)
_REL_ATTRS = tuple(qn(a) for a in ("r:embed", "r:link", "r:id"))


def _find_placeholder(slide, types, skip=0):
    """slide에서 주어진 형식의 자리표시자를 idx 순서로 찾는다."""
    found = sorted(
        (ph for ph in slide.placeholders if ph.placeholder_format.type in types),
        key=lambda ph: ph.placeholder_format.idx,
    )
    return found[skip] if len(found) > skip else None


def _apply_runs(paragraph, runs: List[Run], font_size=None):
    for r in runs:
        run = paragraph.add_run()
        run.text = r.text
        run.font.bold = r.bold or None
        run.font.italic = r.italic or None
        if r.code:
            run.font.name = _CODE_FONT
        if font_size is not None:
            run.font.size = font_size


def _set_textbox_bullet(p, para: Paragraph):
    """일반 텍스트 상자 문단에 글머리 기호/번호 XML을 설정한다.

    자리표시자는 레이아웃/마스터의 목록 스타일을 상속받지만, 텍스트 상자는
    명시적으로 a:buChar / a:buAutoNum 을 넣어야 글머리 기호가 보인다.
    """
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if para.bullet or para.numbered:
        pPr.set("marL", str(int(_INDENT_PER_LEVEL) * (para.level + 1)))
        pPr.set("indent", str(-int(_INDENT_PER_LEVEL)))
        if para.bullet:
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        else:
            pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _fill_text_frame(text_frame, paragraphs: List[Paragraph], textbox_mode=False):
    """텍스트 프레임에 문단들을 채운다. textbox_mode면 글머리 기호 XML도 직접 설정."""
    text_frame.clear()
    for i, para in enumerate(paragraphs):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.level = para.level
        _apply_runs(p, para.runs, font_size=_TEXTBOX_FONT_SIZE if textbox_mode else None)
        if textbox_mode:
            _set_textbox_bullet(p, para)


def _normalized_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return " ".join(shape.text_frame.text.split())


def _replace_shape_text(shape, lines: List[str]):
    """도형의 텍스트를 서식(글꼴/크기/색)을 보존한 채 lines로 교체한다.

    첫 문단/첫 런의 서식을 원형으로 삼아 줄 수만큼 문단을 만든다.
    """
    txBody = shape.text_frame._txBody
    paragraphs = txBody.findall(qn("a:p"))
    if not paragraphs:
        shape.text_frame.text = "\n".join(lines)
        return
    proto = copy.deepcopy(paragraphs[0])
    for p in paragraphs:
        txBody.remove(p)
    for line in lines or [""]:
        p = copy.deepcopy(proto)
        runs = p.findall(qn("a:r"))
        for el in runs[1:] + p.findall(qn("a:br")):
            p.remove(el)
        if runs:
            t = runs[0].find(qn("a:t"))
            t.text = line
        else:
            r = p.makeelement(qn("a:r"), {})
            t = r.makeelement(qn("a:t"), {})
            t.text = line
            r.append(t)
            p.append(r)
        txBody.append(p)


class DeckBuilder:
    def __init__(self, deck: Deck, template_path: Optional[Path] = None,
                 base_dir: Optional[Path] = None):
        self._deck = deck
        self._template_path = template_path
        self._base_dir = base_dir or Path.cwd()
        self._prs = (
            Presentation(str(template_path)) if template_path else Presentation()
        )
        self._resolver = LayoutResolver(self._prs, template_path)
        # 템플릿의 예시 슬라이드: 프로토타입 원본으로 쓰고 빌드 후 제거한다
        self._template_slides = list(self._prs.slides)
        self._template_slide_ids = list(self._prs.slides._sldIdLst)

    def build(self) -> Presentation:
        for spec in self._deck.slides:
            self._add_slide(spec)
        self._strip_template_slides()
        return self._prs

    def _strip_template_slides(self):
        for sldId in self._template_slide_ids:
            rId = sldId.get(qn("r:id"))
            self._prs.part.drop_rel(rId)
            self._prs.slides._sldIdLst.remove(sldId)

    # --- 슬라이드 생성 ---------------------------------------------------

    def _add_slide(self, spec: Slide):
        layout_name = spec.layout or self._auto_layout(spec)
        target = self._resolver.get(layout_name)
        if isinstance(target, PrototypeSpec):
            self._add_from_prototype(spec, target, layout_name)
            return

        slide = self._prs.slides.add_slide(target)
        title_ph = None
        if spec.title is not None:
            title_ph = _find_placeholder(slide, _TITLE_TYPES)
            if title_ph is not None:
                title_ph.text_frame.text = spec.title
            else:
                self._add_title_textbox(slide, spec.title)
        if spec.eyebrow:
            self._add_eyebrow(
                slide, spec.eyebrow, title_ph, center=layout_name in ("title", "section")
            )

        if layout_name == "title":
            self._fill_title_slide(slide, spec)
        elif spec.left or spec.right:
            self._fill_body(slide, spec.left, body_skip=0)
            self._fill_body(slide, spec.right, body_skip=1)
            if spec.blocks:
                self._place_floating_blocks(slide, spec.blocks)
        else:
            self._fill_body(slide, spec.blocks, body_skip=0)

        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    @staticmethod
    def _auto_layout(spec: Slide) -> str:
        if spec.left or spec.right:
            return "two-content"
        if spec.title and not spec.blocks:
            return "title-only"
        if not spec.title and not spec.blocks:
            return "blank"
        return "content"

    # --- 프로토타입 방식 ---------------------------------------------------

    def _add_from_prototype(self, spec: Slide, proto: PrototypeSpec, layout_name: str):
        if not 0 <= proto.slide_index < len(self._template_slides):
            raise ValueError(
                f"프로토타입 슬라이드 인덱스 범위 초과: {proto.slide_index} "
                f"(템플릿 슬라이드 {len(self._template_slides)}장)"
            )
        source = self._template_slides[proto.slide_index]
        slide = self._clone_slide(source)
        fields = self._field_lines(spec, layout_name)

        for shape in list(slide.shapes):
            text = _normalized_text(shape)
            if not text:
                continue
            if any(marker in text for marker in proto.remove):
                shape._element.getparent().remove(shape._element)
                continue
            for marker, field_name in proto.replace.items():
                if marker in text:
                    _replace_shape_text(shape, fields.get(field_name, [""]))
                    break

        area = self._resolve_area(proto.body)
        if layout_name != "title":
            if spec.left or spec.right:
                left_area, right_area = self._split_area(area)
                self._fill_area(slide, spec.left, left_area)
                self._fill_area(slide, spec.right, right_area)
                if spec.blocks:
                    self._fill_area(slide, spec.blocks, area)
            else:
                self._fill_area(slide, spec.blocks, area)

        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    def _field_lines(self, spec: Slide, layout_name: str) -> Dict[str, List[str]]:
        meta = self._deck.meta
        subtitle_lines: List[str] = []
        if layout_name == "title":
            # 표지: 파서가 blocks에 넣어둔 부제/작성자/날짜 문단 사용
            subtitle_lines = [b.text for b in spec.blocks if isinstance(b, Paragraph)]
        return {
            "title": [spec.title] if spec.title else [""],
            "subtitle": subtitle_lines or [""],
            "project": [meta.get("title", "")],
            "author": [meta.get("author", "")],
            "date": [meta.get("date", "")],
        }

    def _clone_slide(self, source):
        slide = self._prs.slides.add_slide(source.slide_layout)
        spTree = slide.shapes._spTree
        # add_slide가 레이아웃에서 복제해 넣은 자리표시자 제거 (원본 슬라이드 것 사용)
        for shape in list(slide.shapes):
            spTree.remove(shape._element)
        for child in source.shapes._spTree:
            if child.tag not in _CLONABLE_TAGS:
                continue
            el = copy.deepcopy(child)
            self._rewire_rels(source.part, slide.part, el)
            spTree.append(el)
        return slide

    @staticmethod
    def _rewire_rels(src_part, dst_part, element):
        """복제된 XML이 참조하는 관계(그림/하이퍼링크)를 새 슬라이드 파트로 옮긴다."""
        for node in element.iter():
            for attr in _REL_ATTRS:
                rid = node.get(attr)
                if not rid:
                    continue
                rel = src_part.rels[rid]
                if rel.is_external:
                    new_rid = dst_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
                node.set(attr, new_rid)

    def _resolve_area(self, body: Optional[List[float]]):
        if body is None:
            return self._default_area()
        sw, sh = self._prs.slide_width, self._prs.slide_height
        scale = (sw, sh, sw, sh)
        return tuple(
            Emu(int(v * s)) if v <= 1.0 else Emu(int(v))
            for v, s in zip(body, scale)
        )

    @staticmethod
    def _split_area(area, gutter_ratio=0.04):
        left, top, width, height = area
        gutter = Emu(int(width * gutter_ratio))
        col = Emu(int((width - gutter) / 2))
        return (
            (left, top, col, height),
            (Emu(int(left + col + gutter)), top, col, height),
        )

    # --- 본문 채우기 -------------------------------------------------------

    def _fill_title_slide(self, slide, spec: Slide):
        subtitle = _find_placeholder(slide, (PP_PLACEHOLDER.SUBTITLE,)) or _find_placeholder(
            slide, _BODY_TYPES
        )
        paragraphs = [b for b in spec.blocks if isinstance(b, Paragraph)]
        if subtitle is not None and paragraphs:
            _fill_text_frame(subtitle.text_frame, paragraphs)

    def _fill_body(self, slide, blocks: List[Block], body_skip: int):
        if not blocks:
            return
        body = _find_placeholder(slide, _BODY_TYPES, skip=body_skip)
        if body is None:
            self._fill_area(slide, blocks, self._default_area())
            return
        text_blocks = [b for b in blocks if isinstance(b, Paragraph)]
        rich_blocks = [b for b in blocks if not isinstance(b, Paragraph)]
        if text_blocks:
            _fill_text_frame(body.text_frame, text_blocks)
        if rich_blocks:
            area = self._block_area(body, below_text=bool(text_blocks))
            self._place_rich_blocks(slide, rich_blocks, area)

    def _fill_area(self, slide, blocks: List[Block], area):
        """자리표시자 없이 주어진 영역에 텍스트 상자/표/그림으로 블록을 배치한다."""
        if not blocks:
            return
        left, top, width, height = area
        text_blocks = [b for b in blocks if isinstance(b, Paragraph)]
        rich_blocks = [b for b in blocks if not isinstance(b, Paragraph)]
        if text_blocks:
            text_height = height if not rich_blocks else Emu(int(height * 0.5))
            box = slide.shapes.add_textbox(left, top, width, text_height)
            box.text_frame.word_wrap = True
            _fill_text_frame(box.text_frame, text_blocks, textbox_mode=True)
        if rich_blocks:
            rich_top = top if not text_blocks else Emu(int(top + height * 0.55))
            rich_height = height if not text_blocks else Emu(int(height * 0.45))
            self._place_rich_blocks(slide, rich_blocks, (left, rich_top, width, rich_height))

    def _add_eyebrow(self, slide, text: str, title_ph, center=False):
        """제목 위의 포인트 컬러 강조 라벨. 색은 템플릿 테마의 강조1을 따른다."""
        from pptx.enum.text import PP_ALIGN

        sw, sh = self._prs.slide_width, self._prs.slide_height
        if title_ph is not None:
            left, width = title_ph.left, title_ph.width
            # 센터형(표지/간지)은 인용 부호 장식과 겹치지 않게 더 위로
            offset = 0.105 if center else 0.05
            top = Emu(max(int(title_ph.top - Emu(int(sh * offset))), 0))
        else:
            left, width = Emu(int(sw * 0.04)), Emu(int(sw * 0.92))
            top = Emu(int(sh * 0.02))
        box = slide.shapes.add_textbox(left, top, width, Emu(int(sh * 0.045)))
        tf = box.text_frame
        tf.text = text
        p = tf.paragraphs[0]
        if center:
            p.alignment = PP_ALIGN.CENTER
        for run in p.runs:
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            run.font._rPr.set("spc", "300")  # 자간 3pt

    def _add_title_textbox(self, slide, title: str):
        """제목 자리표시자가 없는 레이아웃을 위한 대비책."""
        sw, sh = self._prs.slide_width, self._prs.slide_height
        box = slide.shapes.add_textbox(
            Emu(int(sw * 0.04)), Emu(int(sh * 0.03)),
            Emu(int(sw * 0.92)), Emu(int(sh * 0.1)),
        )
        tf = box.text_frame
        tf.text = title
        for run in tf.paragraphs[0].runs:
            run.font.size = Pt(20)
            run.font.bold = True

    # --- 표/그림 배치 -----------------------------------------------------

    def _default_area(self):
        sw, sh = self._prs.slide_width, self._prs.slide_height
        margin = Emu(int(sw * 0.06))
        top = Emu(int(sh * 0.22))
        return margin, top, Emu(sw - 2 * margin), Emu(sh - top - margin)

    def _block_area(self, body, below_text: bool):
        left, top = body.left, body.top
        width, height = body.width, body.height
        if below_text:
            # 텍스트가 위쪽 절반을 차지한다고 보고 아래 절반에 배치
            top = Emu(int(top + height * 0.55))
            height = Emu(int(height * 0.45))
        return left, top, width, height

    def _place_floating_blocks(self, slide, blocks: List[Block]):
        self._place_rich_blocks(
            slide,
            [b for b in blocks if not isinstance(b, Paragraph)],
            self._default_area(),
        )

    def _place_rich_blocks(self, slide, blocks: List[Block], area):
        left, top, width, height = area
        cursor = top
        per_block = Emu(int(height / max(len(blocks), 1)))
        for block in blocks:
            if isinstance(block, Table):
                self._add_table(slide, block, left, cursor, width)
            elif isinstance(block, Image):
                self._add_image(slide, block, left, cursor, width, per_block)
            elif isinstance(block, CodeBlock):
                self._add_code_window(slide, block, left, cursor, width, per_block)
            elif isinstance(block, CardGrid):
                self._add_card_grid(slide, block, left, cursor, width, per_block)
            elif isinstance(block, Bar):
                self._add_bar(slide, block, left, cursor, width, per_block)
            cursor = Emu(int(cursor + per_block))

    def _add_table(self, slide, table: Table, left, top, width):
        n_rows = len(table.rows)
        n_cols = max(len(r) for r in table.rows)
        shape = slide.shapes.add_table(
            n_rows, n_cols, left, top, width, Emu(Pt(24) * n_rows)
        )
        for ri, row in enumerate(table.rows):
            for ci in range(n_cols):
                cell = shape.table.cell(ri, ci)
                runs = row[ci] if ci < len(row) else [Run("")]
                tf = cell.text_frame
                tf.clear()
                _apply_runs(tf.paragraphs[0], runs)
                if ri == 0 and table.has_header:
                    for run in tf.paragraphs[0].runs:
                        run.font.bold = True

    # --- 카드/강조 바 -------------------------------------------------------

    def _add_round_card(self, slide, left, top, width, height, tinted: bool):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.adjustments[0] = 0.12
        card.fill.solid()
        card.shadow.inherit = False
        if tinted:
            # 포인트 컬러를 옅게 깐 강조 카드: 강조1 + 투명도
            card.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            clr = card._element.spPr.find(qn("a:solidFill"))[0]
            clr.append(clr.makeelement(qn("a:alpha"), {"val": "10000"}))
            card.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            card.line.width = Pt(1)
        else:
            card.fill.fore_color.rgb = _CARD_GRAY
            card.line.color.rgb = _CARD_LINE
            card.line.width = Pt(0.75)
        return card

    @staticmethod
    def _styled_para(tf, first: bool, text: str, *, size, bold=False, display=False,
                     accent=False, ink=False, color=None, spacing=None, align=None):
        from pptx.enum.text import PP_ALIGN

        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if align == "center":
            p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold or None
        if display:
            run.font.name = _DISP_FONT
        if accent:
            run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        elif ink:
            run.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
        elif color is not None:
            run.font.color.rgb = color
        if spacing:
            run.font._rPr.set("spc", str(spacing))
        return p

    def _add_card_grid(self, slide, grid: CardGrid, left, top, width, max_height):
        n = len(grid.cards)
        if n == 0:
            return
        cols = 2 if n >= 4 else n
        rows = -(-n // cols)
        is_pair = n == 2 and not grid.numbered
        gap_x = Emu(int(Inches(0.9) if is_pair else Inches(0.28)))
        gap_y = Emu(int(Inches(0.27)))
        card_w = Emu(int((width - (cols - 1) * gap_x) / cols))
        tall = grid.numbered or any(c.label or c.sub for c in grid.cards)
        max_card_h = Inches(1.7) if tall else Inches(0.8)
        card_h = Emu(min(int(max_card_h), int((max_height - (rows - 1) * gap_y) / rows)))

        for i, card in enumerate(grid.cards):
            col, row = i % cols, i // cols
            x = Emu(int(left + col * (card_w + gap_x)))
            y = Emu(int(top + row * (card_h + gap_y)))
            self._add_round_card(slide, x, y, card_w, card_h, card.tinted)
            if grid.numbered:
                self._fill_numbered_card(slide, card, i + 1, x, y, card_w, card_h)
            else:
                self._fill_stack_card(slide, card, x, y, card_w, card_h)

        if is_pair:
            # 비교 카드 사이 화살표
            arrow = slide.shapes.add_textbox(
                Emu(int(left + card_w)), Emu(int(top + (card_h - Inches(0.6)) / 2)),
                gap_x, Emu(int(Inches(0.6))),
            )
            arrow.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._styled_para(arrow.text_frame, True, "→",
                              size=28, color=_ARROW_GRAY, align="center")

    def _fill_stack_card(self, slide, card: Card, x, y, w, h):
        """가운데 정렬 스택 카드: 라벨 / 큰 텍스트 / 보조 설명."""
        box = slide.shapes.add_textbox(
            Emu(int(x + Inches(0.15))), y, Emu(int(w - Inches(0.3))), h
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        if card.label:
            self._styled_para(
                tf, first, card.label, size=10.5, bold=True, spacing=250,
                align="center", accent=card.tinted,
                color=None if card.tinted else _CARD_MUTED,
            )
            first = False
        if card.big:
            self._styled_para(
                tf, first, card.big, size=21, bold=True, display=True,
                align="center", accent=card.tinted, ink=not card.tinted,
            )
            first = False
        if card.sub:
            self._styled_para(tf, first, card.sub, size=12.5,
                              color=_CARD_SUB, align="center")

    def _fill_numbered_card(self, slide, card: Card, number: int, x, y, w, h):
        """번호 원형 배지 + 좌측 정렬 텍스트 카드."""
        badge_d = Emu(int(Inches(0.3)))
        badge_y = Emu(int(y + Inches(0.25))) if card.sub else Emu(int(y + (h - badge_d) / 2))
        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Emu(int(x + Inches(0.25))), badge_y, badge_d, badge_d
        )
        badge.fill.solid()
        badge.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        badge.line.fill.background()
        badge.shadow.inherit = False
        btf = badge.text_frame
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
        self._styled_para(btf, True, str(number), size=11.5, bold=True,
                          color=RGBColor.from_string("FFFFFF"), align="center")

        box = slide.shapes.add_textbox(
            Emu(int(x + Inches(0.68))), y, Emu(int(w - Inches(0.85))), h
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        self._styled_para(tf, True, card.big, size=15 if card.sub else 13.5,
                          bold=True, display=bool(card.sub),
                          accent=card.tinted, ink=not card.tinted)
        if card.sub:
            self._styled_para(tf, False, card.sub, size=11.5, color=_CARD_SUB)

    def _add_bar(self, slide, bar: Bar, left, top, width, max_height):
        """가로 강조 바: 큰 단어 + 작은 라벨 + 본문 (gen_deck.js의 처방 바)."""
        height = Emu(min(int(Inches(0.92)), int(max_height)))
        self._add_round_card(slide, left, top, width, height, tinted=True)

        # 단어 폭 추정 (CJK는 넓게)
        wide = sum(1 for ch in bar.word if ord(ch) > 0x2E80)
        word_w = Emu(int(Inches(min(0.4 + wide * 0.3 + (len(bar.word) - wide) * 0.16, 3.5))))
        if bar.word:
            box = slide.shapes.add_textbox(Emu(int(left + Inches(0.3))), top, word_w, height)
            box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._styled_para(box.text_frame, True, bar.word,
                              size=20, bold=True, display=True, accent=True)

        text_x = Emu(int(left + Inches(0.4) + (word_w if bar.word else 0)))
        box = slide.shapes.add_textbox(
            text_x, top, Emu(int(left + width - text_x - Inches(0.3))), height
        )
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        first = True
        if bar.label:
            self._styled_para(tf, first, bar.label, size=9.5, bold=True,
                              accent=True, spacing=250)
            first = False
        if bar.text:
            self._styled_para(tf, first, bar.text, size=13.5, bold=True, ink=True)

    def _add_code_window(self, slide, code: CodeBlock, left, top, width, max_height):
        """다크 코드 윈도우: 라운드 카드 + 헤더(신호등 점, 파일명) + 모노 코드."""
        head_h = Emu(int(Pt(28)))
        body_h = Emu(int(Pt(20) * max(len(code.lines), 1) + Pt(24)))
        height = Emu(min(int(head_h + body_h), int(max_height)))

        win = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        win.adjustments[0] = 0.06
        win.fill.solid()
        win.fill.fore_color.rgb = _CODE_BG
        win.line.fill.background()
        win.shadow.inherit = False

        head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, head_h)
        head.adjustments[0] = 0.5
        head.fill.solid()
        head.fill.fore_color.rgb = _CODE_HEAD
        head.line.fill.background()
        head.shadow.inherit = False

        dot_d = Emu(int(Pt(7)))
        for i, color in enumerate(_CODE_DOTS):
            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Emu(int(left + Pt(14) + i * Pt(13))),
                Emu(int(top + (head_h - dot_d) / 2)),
                dot_d, dot_d,
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = RGBColor.from_string(color)
            dot.line.fill.background()
            dot.shadow.inherit = False

        if code.label:
            label = slide.shapes.add_textbox(
                Emu(int(left + Pt(56))), top, Emu(int(width - Pt(70))), head_h
            )
            label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            label.text_frame.text = code.label
            for run in label.text_frame.paragraphs[0].runs:
                run.font.name = _CODE_MONO
                run.font.size = Pt(10)
                run.font.color.rgb = _CODE_DIM

        body = slide.shapes.add_textbox(
            Emu(int(left + Pt(16))), Emu(int(top + head_h + Pt(8))),
            Emu(int(width - Pt(32))), Emu(int(height - head_h - Pt(16))),
        )
        tf = body.text_frame
        tf.word_wrap = True
        for i, line in enumerate(code.lines or [""]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.name = _CODE_MONO
            run.font.size = Pt(12.5)
            run.font.color.rgb = _CODE_TXT

    def _add_image(self, slide, image: Image, left, top, width, max_height):
        path = Path(image.path)
        if not path.is_absolute():
            path = self._base_dir / path
        if not path.exists():
            raise FileNotFoundError(f"이미지 파일을 찾을 수 없습니다: {image.path}")
        pic = slide.shapes.add_picture(str(path), left, top, width=width)
        if pic.height > max_height:
            # 영역을 넘으면 높이에 맞춰 비율 유지 축소
            scale = max_height / pic.height
            pic.width = Emu(int(pic.width * scale))
            pic.height = Emu(int(pic.height * scale))


def build_deck(deck: Deck, template_path: Optional[Path] = None,
               base_dir: Optional[Path] = None) -> Presentation:
    """Deck 모델을 PPTX Presentation으로 빌드한다."""
    return DeckBuilder(deck, template_path, base_dir).build()
