"""PPTX 템플릿의 슬라이드 레이아웃을 시맨틱 이름으로 해석한다.

어떤 템플릿을 끼우든 동작하도록 다음 순서로 레이아웃을 찾는다.

1. 사이드카 매핑 파일: 템플릿 옆의 ``<템플릿이름>.layouts.json``
   (예: ``my.pptx.layouts.json``) 에서 시맨틱 이름 → 레이아웃 이름/인덱스
2. 레이아웃 이름 키워드 매칭 (영문/한글 기본 이름)
3. 자리표시자(placeholder) 구성 휴리스틱
4. 표준 Office 템플릿 인덱스 (0=제목, 1=제목 및 내용, ...)

사이드카 값이 JSON 객체이면 **프로토타입 방식**으로 동작한다. 디자인이
레이아웃 자리표시자가 아니라 예시 슬라이드의 텍스트 상자에 들어있는
기업 템플릿을 위해, 템플릿의 예시 슬라이드를 통째로 복제한 뒤 마커
텍스트를 치환한다::

    {
      "title": {
        "prototype": 0,
        "replace": {"프로젝트(솔루션)": "title", "문서(기능)": "subtitle"},
        "remove": ["표지타입", "필요 시 고객사 로고"]
      },
      "content": {
        "prototype": 3,
        "replace": {"속지별 타이틀": "title"},
        "body": [0.05, 0.14, 0.9, 0.78]
      }
    }

- ``prototype``: 복제할 템플릿 예시 슬라이드 인덱스(0부터)
- ``replace``: 마커 부분 문자열 → 채울 필드
  (``title`` / ``subtitle`` / ``project`` / ``author`` / ``date``)
- ``remove``: 해당 마커가 포함된 도형을 삭제
- ``body``: 본문 블록을 배치할 영역 ``[left, top, width, height]``
  (1.0 이하 값은 슬라이드 크기 대비 비율, 초과 값은 EMU)
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

REPLACE_FIELDS = ("title", "subtitle", "project", "author", "date")


@dataclass
class PrototypeSpec:
    """예시 슬라이드 복제 기반 레이아웃 정의."""

    slide_index: int
    replace: Dict[str, str] = field(default_factory=dict)  # 마커 → 필드
    remove: List[str] = field(default_factory=list)
    body: Optional[List[float]] = None

    @classmethod
    def from_dict(cls, semantic: str, data: dict) -> "PrototypeSpec":
        if "prototype" not in data:
            raise ValueError(f"'{semantic}': 프로토타입 객체에는 'prototype' 키가 필요합니다")
        replace = {str(k): str(v) for k, v in data.get("replace", {}).items()}
        for marker, fld in replace.items():
            if fld not in REPLACE_FIELDS:
                raise ValueError(
                    f"'{semantic}': 알 수 없는 필드 '{fld}' (마커 {marker!r}). "
                    f"사용 가능: {REPLACE_FIELDS}"
                )
        body = data.get("body")
        if body is not None and len(body) != 4:
            raise ValueError(f"'{semantic}': body는 [left, top, width, height] 4개여야 합니다")
        return cls(
            slide_index=int(data["prototype"]),
            replace=replace,
            remove=[str(m) for m in data.get("remove", [])],
            body=body,
        )

SEMANTIC_NAMES = ("title", "section", "content", "two-content", "title-only", "blank")

# 레이아웃 이름에 포함되면 매칭되는 키워드 (소문자 비교)
_NAME_KEYWORDS = {
    "title": ("title slide", "제목 슬라이드", "cover", "표지"),
    "section": ("section", "구역", "간지"),
    "two-content": ("two content", "comparison", "콘텐츠 2개", "비교"),
    "content": ("title and content", "제목 및 내용", "content with caption", "내용", "속지", "본문"),
    "title-only": ("title only", "제목만"),
    "blank": ("blank", "빈 화면", "빈 슬라이드"),
}

# python-pptx/Office 기본 템플릿의 표준 인덱스
_STANDARD_INDEX = {
    "title": 0,
    "content": 1,
    "section": 2,
    "two-content": 3,
    "title-only": 5,
    "blank": 6,
}

_BODY_TYPES = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)


def _placeholder_types(layout):
    return [ph.placeholder_format.type for ph in layout.placeholders]


def _guess_by_placeholders(layout) -> Optional[str]:
    types = _placeholder_types(layout)
    bodies = sum(1 for t in types if t in _BODY_TYPES)
    has_title = any(
        t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) for t in types
    )
    if PP_PLACEHOLDER.CENTER_TITLE in types:
        return "title"
    if bodies >= 2:
        return "two-content"
    if bodies == 1:
        return "content"
    if has_title and bodies == 0:
        return "title-only"
    if not types:
        return "blank"
    return None


class LayoutResolver:
    """템플릿 하나에 대한 시맨틱 이름 → SlideLayout 매핑."""

    def __init__(self, presentation: Presentation, template_path: Optional[Path] = None):
        self._prs = presentation
        self._layouts = [
            layout for master in presentation.slide_masters for layout in master.slide_layouts
        ]
        self._map: Dict[str, object] = {}
        sidecar = self._load_sidecar(template_path)
        for name in SEMANTIC_NAMES:
            layout = self._resolve(name, sidecar)
            if layout is not None:
                self._map[name] = layout

    @staticmethod
    def _load_sidecar(template_path: Optional[Path]) -> Dict[str, object]:
        if template_path is None:
            return {}
        sidecar = Path(str(template_path) + ".layouts.json")
        if not sidecar.exists():
            return {}
        data = json.loads(sidecar.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError(f"{sidecar}: 객체(JSON object) 형식이어야 합니다")
        return {str(k).lower(): v for k, v in data.items()}

    def _by_name(self, target: str):
        for layout in self._layouts:
            if (layout.name or "").strip().lower() == target.strip().lower():
                return layout
        return None

    def _resolve(self, name: str, sidecar: Dict[str, object]):
        # 1. 사이드카 매핑 (프로토타입 객체 / 이름 / 인덱스)
        if name in sidecar:
            ref = sidecar[name]
            if isinstance(ref, dict):
                return PrototypeSpec.from_dict(name, ref)
            if isinstance(ref, int):
                if 0 <= ref < len(self._layouts):
                    return self._layouts[ref]
                raise ValueError(f"레이아웃 인덱스 범위 초과: {name} -> {ref}")
            layout = self._by_name(str(ref))
            if layout is None:
                raise ValueError(f"템플릿에 '{ref}' 레이아웃이 없습니다 ({name})")
            return layout
        # 2. 레이아웃 이름 키워드
        for layout in self._layouts:
            layout_name = (layout.name or "").lower()
            if any(kw in layout_name for kw in _NAME_KEYWORDS[name]):
                return layout
        # 3. 자리표시자 구성 휴리스틱
        for layout in self._layouts:
            if _guess_by_placeholders(layout) == name:
                return layout
        # 4. 표준 인덱스 — 표준 Office 구성(레이아웃 7개 이상)일 때만 신뢰
        if len(self._layouts) >= 7:
            idx = _STANDARD_INDEX[name]
            if idx < len(self._layouts):
                return self._layouts[idx]
        return None

    def get(self, name: str):
        """시맨틱 이름으로 레이아웃(또는 PrototypeSpec)을 얻는다.

        없으면 content → 첫 레이아웃 순서로 대체한다.
        """
        if name in self._map:
            return self._map[name]
        if "content" in self._map:
            return self._map["content"]
        return self._layouts[0]

    def describe(self) -> str:
        """`officedocs layouts` 명령용: 템플릿 레이아웃 목록을 사람이 읽기 좋게 정리."""
        reverse = {}
        prototypes = {}
        for semantic, target in self._map.items():
            if isinstance(target, PrototypeSpec):
                prototypes[semantic] = target
            else:
                reverse[id(target)] = semantic
        lines = []
        for i, layout in enumerate(self._layouts):
            semantic = reverse.get(id(layout), "-")
            phs = ", ".join(
                f"{ph.placeholder_format.idx}:{str(ph.placeholder_format.type).split('.')[-1].split(' ')[0]}"
                for ph in layout.placeholders
            )
            lines.append(f"[{i}] {layout.name!r}  semantic={semantic}  placeholders=({phs})")
        for semantic, spec in sorted(prototypes.items()):
            markers = ", ".join(f"{m!r}→{f}" for m, f in spec.replace.items())
            lines.append(
                f"(prototype) semantic={semantic}  slide={spec.slide_index}  replace=({markers})"
            )
        return "\n".join(lines)
