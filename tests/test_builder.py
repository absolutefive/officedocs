import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from officedocs.builder import build_deck
from officedocs.parser import parse_document

ROOT = Path(__file__).resolve().parent.parent
EXAMPLE = ROOT / "examples" / "sample.md"
TEMPLATES = ROOT / "templates"


def _build(template=None):
    deck = parse_document(EXAMPLE.read_text(encoding="utf-8"))
    prs = build_deck(deck, template, base_dir=EXAMPLE.parent)
    # 저장 → 재오픈 라운드트립으로 파일 유효성까지 확인
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return Presentation(buf)


def test_build_with_builtin_default():
    prs = _build()
    assert len(prs.slides) == 9  # 표지 1 + 본문 8


def test_titles_are_native_text():
    prs = _build()
    first = prs.slides[0]
    titles = [s for s in first.shapes if s.has_text_frame and "officedocs" in s.text_frame.text]
    assert titles, "표지 제목이 텍스트 개체로 존재해야 한다"


def test_table_is_native_object():
    prs = _build()
    tables = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_table
    ]
    assert tables, "표가 네이티브 표 개체여야 한다"
    assert tables[0].table.cell(0, 0).text == "기능"


def test_image_is_picture_shape():
    prs = _build()
    pictures = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert pictures, "그림이 Picture 개체여야 한다"


def test_notes_round_trip():
    prs = _build()
    notes = [
        slide.notes_slide.notes_text_frame.text
        for slide in prs.slides
        if slide.has_notes_slide
    ]
    assert any("도입부" in n for n in notes)


def test_no_raster_only_output():
    """본문 텍스트가 이미지가 아니라 실제 텍스트 프레임에 존재하는지 확인."""
    prs = _build()
    all_text = "\n".join(
        shape.text_frame.text
        for slide in prs.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "내용은 마크다운 텍스트로 작성" in all_text


@pytest.mark.parametrize("name", ["default.pptx", "dark.pptx"])
def test_build_with_bundled_templates(name):
    template = TEMPLATES / name
    if not template.exists():
        pytest.skip("번들 템플릿이 아직 생성되지 않음 (scripts/make_templates.py)")
    prs = _build(template)
    assert len(prs.slides) == 9


def test_missing_image_raises(tmp_path):
    doc = tmp_path / "doc.md"
    doc.write_text("# x\n![](no-such.png)", encoding="utf-8")
    deck = parse_document(doc.read_text(encoding="utf-8"))
    with pytest.raises(FileNotFoundError):
        build_deck(deck, base_dir=tmp_path)
